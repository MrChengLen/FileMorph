# SPDX-License-Identifier: AGPL-3.0-or-later
"""Deterministic QA fixtures for the "Insert PDF pages" and photo-orientation
manual test round.

Generates the files the manual checklist refers to:

* ``hausarbeit-30-seiten.pdf`` — 30 A4 pages, each with "Seite N" in large
  type, and bookmarks "Kapitel 1" (page 3), "Kapitel 2" (page 12),
  "Kapitel 3" (page 21) and "Anhang" (page 28): the PDF to insert into
* ``anhang-2-seiten.pdf`` — two US-Letter pages, "Anhang A" and "Anhang B"
* ``erklaerung.docx`` — a one-page Word document, "Eidesstattliche Erklärung"
* ``unterschrift-hochformat.jpg`` — a portrait photo as a phone stores it:
  the pixels lie landscape, EXIF orientation 6 turns them upright; "OBEN"
  marks the top
* ``foto-hochformat.jpg`` — the same kind of phone photo, for converting and
  compressing (it must come back upright, not lying on its side)
* ``palette-bild.gif`` — a colourful palette image (every GIF is one); converted
  to PNG it must keep its colours
* ``notizen.txt`` and ``ausarbeitung.odt`` — files the insertion refuses
* ``gesperrt.pdf`` — ``hausarbeit-30-seiten.pdf`` locked with the password
  ``filemorph`` (via ``make_testdata_pdf_error_messages._locked``)

Same call => byte-identical output on every run: the PDFs are assembled object
by object, the images are drawn without randomness, and the Word file is
re-zipped with fixed timestamps. Output directory defaults to
``docs-internal/testdata/pdf-insert/`` next to this repo checkout (the folder
is gitignored — commit this script, never the data) and can be overridden as
the first CLI argument.

Run:
    python scripts/make_testdata_pdf_insert.py [output_dir]
"""

from __future__ import annotations

import io
import sys
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader, PdfWriter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from make_testdata_pdf_error_messages import _locked  # noqa: E402

A4 = (595, 842)
LETTER = (612, 792)
HELVETICA = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>"


def _pdf(pages: list[tuple[tuple[int, int], bytes]]) -> bytes:
    """A PDF with one page per (size, text) entry, the text large and centred-ish."""
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"", HELVETICA]
    kids = []
    for (width, height), text in pages:
        content = b"BT /F1 48 Tf 72 %d Td (%s) Tj ET" % (height // 2, text)
        objects.append(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content))
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 %d %d] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>"
            % (width, height, len(objects))
        )
        kids.append(b"%d 0 R" % len(objects))
    objects[1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (b" ".join(kids), len(kids))
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (number, body)
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    return bytes(out)


def _thesis() -> bytes:
    raw = _pdf([(A4, b"Seite %d" % n) for n in range(1, 31)])
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(raw)))
    for title, page in (("Kapitel 1", 3), ("Kapitel 2", 12), ("Kapitel 3", 21), ("Anhang", 28)):
        writer.add_outline_item(title, page - 1)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _docx() -> bytes:
    """A one-page Word file, re-zipped with fixed timestamps (byte-stable)."""
    from docx import Document

    document = Document()
    document.core_properties.author = "Testdatei"
    document.add_heading("Eidesstattliche Erklärung", level=1)
    document.add_paragraph(
        "Hiermit versichere ich, dass ich die vorliegende Hausarbeit selbstständig "
        "verfasst und keine anderen als die angegebenen Hilfsmittel benutzt habe."
    )
    document.add_paragraph("Ort, Datum: ____________        Unterschrift: ____________")
    buf = io.BytesIO()
    document.save(buf)
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(buf.getvalue())) as src,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst,
    ):
        for info in src.infolist():
            data = src.read(info.filename)
            fixed = zipfile.ZipInfo(info.filename, date_time=(2026, 10, 27, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(fixed, data)
    return out.getvalue()


def _portrait_photo(label: str) -> bytes:
    """A portrait picture (1200 x 1600) stored as a phone does: pixels turned
    to landscape, EXIF orientation 6 to show them upright."""
    upright = Image.new("RGB", (1200, 1600), (235, 235, 225))
    draw = ImageDraw.Draw(upright)
    font = ImageFont.load_default(size=150)
    draw.rectangle((0, 0, 1199, 260), fill=(200, 40, 40))
    draw.text((600, 130), "OBEN", fill=(255, 255, 255), font=font, anchor="mm")
    draw.text(
        (600, 900), label, fill=(30, 30, 30), font=ImageFont.load_default(size=90), anchor="mm"
    )
    draw.line((250, 1250, 950, 1250), fill=(30, 30, 30), width=8)
    stored = upright.transpose(Image.Transpose.ROTATE_90)  # what the sensor saw
    exif = Image.Exif()
    exif[0x0112] = 6  # Orientation: turn 90° clockwise to show
    buf = io.BytesIO()
    stored.save(buf, format="JPEG", quality=90, exif=exif.tobytes())
    return buf.getvalue()


def _palette_gif() -> bytes:
    """Colour bands in a palette image, the shape of every GIF."""
    img = Image.new("RGB", (480, 240))
    draw = ImageDraw.Draw(img)
    colours = [
        (220, 50, 47),
        (38, 139, 210),
        (133, 153, 0),
        (181, 137, 0),
        (108, 113, 196),
        (42, 161, 152),
    ]
    for i, colour in enumerate(colours):
        draw.rectangle((i * 80, 0, i * 80 + 79, 239), fill=colour)
    buf = io.BytesIO()
    img.convert("P", palette=Image.Palette.ADAPTIVE, colors=16).save(buf, format="GIF")
    return buf.getvalue()


def main() -> None:
    out = (
        Path(sys.argv[1])
        if len(sys.argv) > 1
        else Path(__file__).resolve().parent.parent / "docs-internal" / "testdata" / "pdf-insert"
    )
    out.mkdir(parents=True, exist_ok=True)
    thesis = _thesis()
    files = {
        "hausarbeit-30-seiten.pdf": thesis,
        "anhang-2-seiten.pdf": _pdf([(LETTER, b"Anhang A"), (LETTER, b"Anhang B")]),
        "erklaerung.docx": _docx(),
        "unterschrift-hochformat.jpg": _portrait_photo("Unterschrift"),
        "foto-hochformat.jpg": _portrait_photo("Hochformat"),
        "palette-bild.gif": _palette_gif(),
        "notizen.txt": "Diese Datei kann nicht eingefügt werden.\n".encode(),
        "ausarbeitung.odt": b"PK\x03\x04 kein echtes ODT, nur der Name zaehlt",
        "gesperrt.pdf": _locked(thesis, "AES-256", "filemorph"),
    }
    for name, data in files.items():
        (out / name).write_bytes(data)
        print(f"{name}: {len(data):,} bytes")
    print(f"-> {out}")


if __name__ == "__main__":
    main()
