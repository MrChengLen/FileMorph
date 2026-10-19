# SPDX-License-Identifier: AGPL-3.0-or-later
"""XLSX → CSV reads cell values only.

Chart sheets are read without their drawings, the active sheet falls back to
the first worksheet when a chart sheet is active, and a workbook without a
worksheet gets a 400.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile

import openpyxl
import PIL.Image
import pytest
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as XlImage

from app.converters.base import InvalidInputError
from app.converters.spreadsheet import XlsxToCsvConverter


def _workbook_with_chartsheet(path, *, active_chartsheet=False, image=False) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for row in (("Monat", "Umsatz"), ("Jan", 10), ("Feb", 12)):
        ws.append(row)
    if image:
        png = io.BytesIO()
        PIL.Image.new("RGB", (4, 4), "red").save(png, "PNG")
        png.seek(0)
        ws.add_image(XlImage(png), "D2")
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    chartsheet = wb.create_chartsheet("Chart")
    chartsheet.add_chart(chart)
    if active_chartsheet:
        wb.active = wb.sheetnames.index("Chart")
    wb.save(path)


def _rewrite(src, dst, edit) -> None:
    """Copy the package ``src`` to ``dst``, passing each part through ``edit``."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w") as zout:
        for name in zin.namelist():
            zout.writestr(name, edit(name, zin.read(name)))


def _rows(path) -> list[list[str]]:
    with path.open(encoding="utf-8", newline="") as f:
        return list(csv.reader(f))


def test_images_in_a_chart_sheet_are_never_decoded(tmp_path, monkeypatch):
    base, src, out = tmp_path / "base.xlsx", tmp_path / "in.xlsx", tmp_path / "out.csv"
    _workbook_with_chartsheet(base, image=True)
    with zipfile.ZipFile(base) as z:
        picture = next(
            n for n in z.namelist() if n.startswith("xl/drawings/drawing") and b"<pic>" in z.read(n)
        )

    def chartsheet_shows_the_picture(name, data):
        if name.startswith("xl/chartsheets/_rels/"):
            return re.sub(rb'Target="[^"]+"', f'Target="/{picture}"'.encode(), data)
        return data

    _rewrite(base, src, chartsheet_shows_the_picture)

    def decoded(*args, **kwargs):
        raise AssertionError("an image was decoded")

    monkeypatch.setattr(PIL.Image, "open", decoded)
    with pytest.raises(AssertionError, match="an image was decoded"):
        openpyxl.load_workbook(str(src), read_only=True)  # the stock loader reaches the picture
    XlsxToCsvConverter().convert(src, out)
    assert _rows(out) == [["Monat", "Umsatz"], ["Jan", "10"], ["Feb", "12"]]


def test_an_active_chart_sheet_exports_the_first_worksheet(tmp_path):
    src, out = tmp_path / "in.xlsx", tmp_path / "out.csv"
    _workbook_with_chartsheet(src, active_chartsheet=True)
    XlsxToCsvConverter().convert(src, out)
    assert _rows(out) == [["Monat", "Umsatz"], ["Jan", "10"], ["Feb", "12"]]


def test_a_workbook_without_a_worksheet_gets_a_clear_error(tmp_path):
    base, src, out = tmp_path / "base.xlsx", tmp_path / "in.xlsx", tmp_path / "out.csv"
    _workbook_with_chartsheet(base)

    def chart_sheet_only(name, data):
        if name == "xl/workbook.xml":
            return re.sub(rb'<sheet [^>]*name="Data"[^>]*/>', b"", data)
        return data

    _rewrite(base, src, chart_sheet_only)
    with pytest.raises(InvalidInputError, match="no worksheet"):
        XlsxToCsvConverter().convert(src, out)
