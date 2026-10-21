# SPDX-License-Identifier: AGPL-3.0-or-later
"""Table conversions stop at the output cap.

XLSX → CSV, JSON → CSV and CSV → JSON take the caller's output cap as
``max_output_bytes`` and stop writing once the output passes it; the convert
routes answer with 413 ``output_cap_exceeded``.
"""

from __future__ import annotations

import csv
import dataclasses
import io
import json
import re
import zipfile
from datetime import datetime

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

from app.converters.base import OutputTooLargeError
from app.converters.spreadsheet import (
    CsvToJsonConverter,
    CsvToXlsxConverter,
    JsonToCsvConverter,
    XlsxToCsvConverter,
    _CappedText,
    _load_cells_only,
    _rows_within,
)
from app.core.quotas import _MB, QUOTAS


def _sheet_spanning(path, last_cell: str) -> None:
    """A sheet with cells in A1 and ``last_cell`` only."""
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    wb.active[last_cell] = 1
    wb.save(path)


def _rows(path) -> list[list[str]]:
    with path.open(encoding="utf-8", newline="") as f:
        return list(csv.reader(f))


def test_xlsx_to_csv_stops_writing_at_the_cap(tmp_path):
    src, out = tmp_path / "in.xlsx", tmp_path / "out.csv"
    _sheet_spanning(src, "Z2000")
    with pytest.raises(OutputTooLargeError):
        XlsxToCsvConverter().convert(src, out, max_output_bytes=10_000)
    assert out.stat().st_size <= 10_000


def test_a_row_too_long_for_the_cap_is_refused_before_it_is_built():
    out = _CappedText(io.StringIO(), 1_000_000)
    rows = _rows_within(iter([("a" * 50_000,) * 1000]), out, longest=50_000)
    with pytest.raises(OutputTooLargeError):
        next(rows)


def test_ordinary_rows_pass_the_row_check_unchanged():
    rows = [("x", 1, None, 'say "hi"'), (2.5, True, "", None)]
    out = _CappedText(io.StringIO(), 10_000)
    assert list(_rows_within(iter(rows), out, longest=8)) == rows


def _shared_text_book(path, text: str, cells: int) -> None:
    """A sheet whose cells all point at one entry of the shared-strings table, the way Excel saves text."""
    wb = openpyxl.Workbook()
    wb.active["A1"] = 0
    wb.save(path)
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    row = "".join(f'<c r="{get_column_letter(i)}1" t="s"><v>0</v></c>' for i in range(1, cells + 1))
    sheet = re.sub(
        rb'<dimension ref="[^"]*"/>',
        f'<dimension ref="A1:{get_column_letter(cells)}1"/>'.encode(),
        parts["xl/worksheets/sheet1.xml"],
    )
    parts["xl/worksheets/sheet1.xml"] = re.sub(
        rb"<sheetData>.*</sheetData>",
        f'<sheetData><row r="1">{row}</row></sheetData>'.encode(),
        sheet,
        flags=re.S,
    )
    parts["xl/sharedStrings.xml"] = (
        f'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t>{text}</t></si></sst>'
    ).encode()
    parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
        b"</Types>",
        b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>',
    )
    parts["xl/_rels/workbook.xml.rels"] = parts["xl/_rels/workbook.xml.rels"].replace(
        b"</Relationships>",
        b'<Relationship Id="rIdShared" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/></Relationships>',
    )
    with zipfile.ZipFile(path, "w") as z:
        for name, data in parts.items():
            z.writestr(name, data)


def test_xlsx_to_csv_refuses_a_row_of_one_long_shared_text_before_building_it(
    tmp_path, monkeypatch
):
    src, out = tmp_path / "in.xlsx", tmp_path / "out.csv"
    _shared_text_book(src, "a" * 20_000, cells=100)
    with src.open("rb") as upload:
        assert len(_load_cells_only(upload).active._shared_strings) == 1  # the text is shared
    written = []
    write = _CappedText.write

    def counting_write(self, text):
        written.append(len(text))
        return write(self, text)

    monkeypatch.setattr(_CappedText, "write", counting_write)
    with pytest.raises(OutputTooLargeError):
        XlsxToCsvConverter().convert(src, out, max_output_bytes=100_000)
    assert written == []  # refused before csv built the row


def test_xlsx_to_csv_with_an_overstated_size_still_converts(tmp_path):
    base, src, out = tmp_path / "base.xlsx", tmp_path / "in.xlsx", tmp_path / "out.csv"
    wb = openpyxl.Workbook()
    for row in (("a", "b"), ("1", "2")):
        wb.active.append(row)
    wb.save(base)
    with zipfile.ZipFile(base) as zin, zipfile.ZipFile(src, "w") as zout:
        for name in zin.namelist():
            data = zin.read(name)
            if name == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:CV1048576"', data)
            zout.writestr(name, data)
    XlsxToCsvConverter().convert(src, out, max_output_bytes=10_000)
    assert [row[:2] for row in _rows(out)] == [["a", "b"], ["1", "2"]]


def test_xlsx_to_csv_writes_the_cell_values_as_before(tmp_path):
    src, out = tmp_path / "in.xlsx", tmp_path / "out.csv"
    wb = openpyxl.Workbook()
    for row in (("x", 0.1, 1e20), (None,), (True, datetime(2026, 1, 5, 8, 30), 'say "hi"\nok')):
        wb.active.append(row)
    wb.save(src)
    XlsxToCsvConverter().convert(src, out, max_output_bytes=10_000)
    expected = io.StringIO(newline="")
    writer = csv.writer(expected)
    for row in openpyxl.load_workbook(src, read_only=True).active.iter_rows(values_only=True):
        writer.writerow(["" if v is None else str(v) for v in row])
    assert out.read_bytes().decode("utf-8") == expected.getvalue()


def test_json_to_csv_with_more_cells_than_the_cap_is_refused_before_writing(tmp_path):
    src, out = tmp_path / "in.json", tmp_path / "out.csv"
    src.write_text(json.dumps([{f"k{i}": "" for i in range(1000)}] + [{}] * 100), encoding="utf-8")
    with pytest.raises(OutputTooLargeError):
        JsonToCsvConverter().convert(src, out, max_output_bytes=50_000)
    assert not out.exists()


def test_json_to_csv_stops_writing_at_the_cap(tmp_path):
    src, out = tmp_path / "in.json", tmp_path / "out.csv"
    src.write_text(json.dumps([{"a": "ü" * 1000}] * 50), encoding="utf-8")
    with pytest.raises(OutputTooLargeError):
        JsonToCsvConverter().convert(src, out, max_output_bytes=10_000)
    assert out.stat().st_size <= 10_000  # bytes, though "ü" is one character


def test_csv_to_json_stops_writing_at_the_cap(tmp_path):
    src, out = tmp_path / "in.csv", tmp_path / "out.json"
    header = ",".join(f"column{i}" for i in range(200))
    src.write_text(header + "\n" + "x\n" * 100, encoding="utf-8")
    with pytest.raises(OutputTooLargeError):
        CsvToJsonConverter().convert(src, out, max_output_bytes=10_000)
    assert out.stat().st_size <= 10_000


@pytest.mark.parametrize(
    "text",
    [
        "a,b\n1,ü\n2,\n3",
        "a,b\n",
        "a,b\n1,2,3\n",  # a field more than the header
        "k\nline\u2028one\n",
        "k\nline\u2029two\n",
        "k\nline\x85three\n",
        pytest.param(
            "k,v\r\n" + "".join(f'{i},"\u20ac{i} \U0001f600\r\n\u00fc"\r\n' for i in range(5_000)),
            id="rows-past-the-read-buffer",
        ),
    ],
)
def test_csv_to_json_writes_what_json_dumps_would(tmp_path, text):
    src, out = tmp_path / "in.csv", tmp_path / "out.json"
    src.write_bytes(text.encode("utf-8"))  # line endings as given, on Windows too
    CsvToJsonConverter().convert(src, out)
    rows = list(csv.DictReader(io.StringIO(text, newline="")))
    assert out.read_text(encoding="utf-8") == json.dumps(rows, indent=2, ensure_ascii=False)


def test_csv_to_xlsx_keeps_every_cell(tmp_path):
    src, out = tmp_path / "in.csv", tmp_path / "out.xlsx"
    src.write_text('a,b\n1,"x, y"\nü,\n', encoding="utf-8")
    CsvToXlsxConverter().convert(src, out)
    ws = openpyxl.load_workbook(out).active
    assert [[c or "" for c in row] for row in ws.iter_rows(values_only=True)] == [
        ["a", "b"],
        ["1", "x, y"],
        ["ü", ""],
    ]


@pytest.fixture
def one_mb_cap(monkeypatch):
    """The test key is anonymous; give that tier a 1 MB output cap."""
    capped = dataclasses.replace(QUOTAS["anonymous"], output_cap_bytes=_MB)
    monkeypatch.setitem(QUOTAS, "anonymous", capped)


def test_convert_route_answers_413_when_a_table_passes_the_cap(
    client, auth_headers, tmp_path, one_mb_cap
):
    src = tmp_path / "in.xlsx"
    _sheet_spanning(src, "Z50000")
    resp = client.post(
        "/api/v1/convert",
        files={"file": ("sheet.xlsx", src.read_bytes(), "application/octet-stream")},
        data={"target_format": "csv"},
        headers=auth_headers,
    )
    assert resp.status_code == 413
    assert resp.headers["X-FileMorph-Error-Code"] == "output_cap_exceeded"
    assert resp.json()["detail"].startswith(
        "Output too large (over the 1 MB cap). Split the table into smaller files"
    )


def test_batch_route_reports_a_table_that_passes_the_cap(
    client, auth_headers, tmp_path, one_mb_cap
):
    src = tmp_path / "in.xlsx"
    _sheet_spanning(src, "Z50000")
    resp = client.post(
        "/api/v1/convert/batch",
        files=[("files", ("sheet.xlsx", src.read_bytes(), "application/octet-stream"))],
        data={"target_formats": ["csv"]},
        headers=auth_headers,
    )
    assert resp.status_code == 422
    assert resp.json()["files"][0]["error_message"] == (
        "Output too large (over the 1 MB cap). Split the table. Register free for a higher cap."
    )
